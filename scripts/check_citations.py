#!/usr/bin/env python3
"""Citation integrity, and the inventory the ES&T numbering pass will need.

Two jobs:

  --check   (default) every in-text citation resolves to exactly one reference
            entry, and every entry is cited at least once. Exit non-zero
            otherwise. This is what caught nothing on 2026-08-24 only because
            C3 (three Table 1 sources missing from the list) and C4
            (Maloszewski 1996 uncited) had just been fixed.

  --order   print the citations in order of first appearance through the
            document -- paragraphs and tables interleaved in true body order,
            which is what ES&T's "numbered in order of appearance" means. This
            is the input to the numbering pass.

The numbering pass itself (C1) is deliberately NOT automated here. It must be
run once the content is frozen: any change that moves text -- relocating
Table 1 to the SI, a rewritten abstract, a new citation in a revision --
renumbers everything downstream of it. Run this with --order at that point and
convert from the printed list.

Usage:  python check_citations.py [--order] [--si]
"""
from __future__ import annotations

import argparse
import os
import re
import sys

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

FC = r"D:\projects\carbon isotopic\DLMPI_reanalysis\From_Claude_chat"
MAIN = os.path.join(FC, "manuscript", "WRR_Groundwater_Age_Identifiability.docx")
SI = os.path.join(FC, "manuscript", "WRR_Groundwater_Age_Identifiability_SI.docx")

# Author-date citations. Covers "Name, 1997", "Name and Other, 1997",
# "Name et al., 1997", the narrative "Name et al. (1997)" form, and the one
# corporate author (IAEA/WMO).
CIT = re.compile(
    r"([A-Z][A-Za-z\u00c0-\u024f'\-]+"
    r"(?:\s+(?:and|&)\s+[A-Z][A-Za-z\u00c0-\u024f'\-]+)?"
    r"(?:\s+et\s+al\.?)?|IAEA/WMO)[,]?\s*\(?((?:19|20)\d{2}[ab]?)\)?")


#: tokens that look like an author but precede a DATE, not a year of
#: publication -- Table S11's benchmark labels and a figure annotation
NOT_AUTHORS = {"Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep",
               "Oct", "Nov", "Dec", "NE", "NM", "CA", "TX", "SSW", "TU"}


def blocks(doc):
    """Paragraphs and tables in true document order."""
    for ch in doc.element.body.iterchildren():
        if ch.tag.endswith('}p'):
            yield Paragraph(ch, doc)
        elif ch.tag.endswith('}tbl'):
            yield Table(ch, doc)


def inventory(path):
    """(ordered citation keys, reference entries) for one document."""
    d = docx.Document(path)
    # The SI has TWO paragraphs reading "References": the Contents entry near
    # the top and the list itself at the end. Stopping at the first one scanned
    # nothing at all, so anchor on the LAST.
    _all = list(d.paragraphs)
    _refhead = None
    for _i, _q in enumerate(_all):
        if _q.text.strip() == "References":
            _refhead = _q._p
    seq, stop = [], False
    for b in blocks(d):
        if isinstance(b, Paragraph):
            if b._p is _refhead:
                stop = True
                continue
            txt = b.text
        else:
            txt = " ; ".join(c.text for r in b.rows for c in r.cells)
        if stop:
            continue
        for m in CIT.finditer(txt):
            name = re.sub(r"\s+", " ", m.group(1)).strip()
            if name in NOT_AUTHORS:
                continue          # "Jun 2007", "Aug 2004", "NE 1963" -- dates
            for yr in [m.group(2)] + re.findall(
                    r"(?:19|20)\d{2}[ab]?",
                    txt[m.end():m.end() + 12].split(")")[0]):
                # a combined citation such as "Musgrove et al. (2023a, 2023b)"
                # names the author once and two years; both must be counted
                key = (name, yr)
                if key not in seq:
                    seq.append(key)
    ps = d.paragraphs
    rs = [i for i, p in enumerate(ps) if p.text.strip() == "References"]
    refs = []
    if rs:
        refs = [(i, ps[i].text.strip())
                for i in range(rs[-1] + 1, len(ps)) if ps[i].text.strip()]
    return seq, refs


def resolve(cit, refs):
    """The reference entries a citation could mean. Exactly one is required."""
    author, year = cit
    first = author.split(" and ")[0].split(" et al")[0].strip()
    second = author.split(" and ")[1].strip() if " and " in author else None
    yr, suf = year.rstrip("ab"), year[len(year.rstrip("ab")):]
    out = []
    for i, t in refs:
        if not t.startswith(first):
            continue
        if suf:
            if (yr + suf) not in t:
                continue
        else:
            # a bare year must not match an a/b-suffixed entry
            if yr not in t or re.search(re.escape(yr) + r"[ab]", t):
                continue
        if second and second not in t:
            continue
        out.append(i)
    return out


def numbered_state(refs):
    """True once the ES&T numbering pass has run: entries begin '(N) '."""
    return bool(refs) and all(re.match(r"^\(\d+\)\s", t) for _, t in refs)


def check_numbered(path, label):
    """Invariants AFTER conversion to numbered superscripts.

    The author-date checks cannot apply once the names are gone, so the
    equivalent guarantees are: the list is numbered 1..N contiguously in
    document order, every number is cited as a superscript somewhere in the
    body, and no author-date citation survives.
    """
    d = docx.Document(path)
    _all = list(d.paragraphs)
    refhead = None
    for q in _all:
        if q.text.strip() == "References":
            refhead = q._p
    ps = d.paragraphs
    st = [i for i, q in enumerate(ps) if q.text.strip() == "References"][-1]
    refs = [ps[i].text.strip() for i in range(st + 1, len(ps))
            if ps[i].text.strip()]
    got = [int(re.match(r"^\((\d+)\)", t).group(1)) for t in refs]
    ok = got == list(range(1, len(refs) + 1))
    print("=== %s (numbered) ===" % label)
    print("  %d references, numbered 1..%d in order: %s"
          % (len(refs), len(refs), ok))

    cited, stop = set(), False
    for ch in d.element.body.iterchildren():
        if ch.tag.endswith('}p'):
            if ch is refhead:
                stop = True
                continue
            paras = [] if stop else [Paragraph(ch, d)]
        elif ch.tag.endswith('}tbl') and not stop:
            paras = [q for row in Table(ch, d).rows for c in row.cells
                     for q in c.paragraphs]
        else:
            continue
        for q in paras:
            for r in q.runs:
                if r.font.superscript and re.fullmatch(r"[\d,–]+",
                                                       r.text or ""):
                    for tok in r.text.split(","):
                        if "–" in tok:
                            a, b = tok.split("–")
                            cited.update(range(int(a), int(b) + 1))
                        else:
                            cited.add(int(tok))
    missing = sorted(set(range(1, len(refs) + 1)) - cited)
    print("  %d of %d numbers cited in the text%s"
          % (len(cited), len(refs),
             "" if not missing else "  MISSING %s" % missing))
    body = chr(10).join(q.text for q in _all[:st])
    left = re.findall(r"\([A-Z][A-Za-z'\-]+[^()]{0,40},\s*(?:19|20)\d{2}",
                      body)
    print("  surviving author-date citations: %d" % len(left))
    good = ok and not missing and not left
    print("  %s" % ("OK" if good else "PROBLEMS ABOVE"))
    return good


def run(path, label, show_order):
    seq, refs = inventory(path)
    if numbered_state(refs):
        return check_numbered(path, label)
    print("=== %s ===" % label)
    print("  %d distinct citations, %d reference entries" % (len(seq), len(refs)))
    if not refs:
        print("  (no reference list found -- skipping)")
        return True
    bad, used = [], set()
    for c in seq:
        hits = resolve(c, refs)
        if len(hits) != 1:
            bad.append((c, len(hits)))
        else:
            used.add(hits[0])
    for c, n in bad:
        print("  UNRESOLVED (%d matches): %s %s" % (n, c[0], c[1]))
    for i, t in refs:
        if i not in used:
            print("  NEVER CITED: %s" % t[:78])
    ok = not bad and len(used) == len(refs)
    print("  %s" % ("OK -- 1:1" if ok else "PROBLEMS ABOVE"))
    if show_order:
        print("\n  order of first appearance (the ES&T numbering):")
        for n, (a, y) in enumerate(seq, 1):
            print("   %3d. %s %s" % (n, a, y))
    return ok


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--order", action="store_true",
                    help="also print the numbering order")
    ap.add_argument("--si", action="store_true", help="check the SI too")
    A = ap.parse_args()
    ok = run(MAIN, "main text", A.order)
    if A.si:
        print()
        ok = run(SI, "supporting information", A.order) and ok
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
