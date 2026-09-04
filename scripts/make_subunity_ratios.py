#!/usr/bin/env python3
"""subunity_ratios.csv and SI Table S10 -- information ratios below unity.

Both were PRE-BISECTION-FIX artifacts. They showed 13 cases with a minimum of
0.816 and "ten of the thirteen" from omitting 14C; recomputed from the current
profile_true_loto.csv the answer is **12 of the 59 defined ratios, minimum
0.7619, eight from 14C** -- which is exactly what the main text has said since
2026-08-23. The fourth review flagged the discrepancy (its item B7) and
proposed changing the main text to match the table; the table was the stale
side, so the table is regenerated instead. Its tracked change must be REJECTED.

The filter matters and has been wrong here twice before: a CENSORED row carries
a ratio value, but it is a bound-limited lower bound, not a measured width
ratio. Only uncensored rows with a finite true-profile ratio count.

Usage:  python make_subunity_ratios.py [--write-docx]
"""
from __future__ import annotations

import argparse
import os
import sys

import numpy as np
import pandas as pd

R = r"D:\projects\carbon isotopic\DLMPI_reanalysis"
LOTO = os.path.join(R, "DLPMI_edwards_handoff", "sweeps_trueprofile",
                    "profile_true_loto.csv")
OUT = os.path.join(R, "From_Claude_chat", "supplementary_data",
                   "subunity_ratios.csv")
SI = os.path.join(R, "From_Claude_chat", "manuscript",
                  "WRR_Groundwater_Age_Identifiability_SI.docx")

PRETTY = {"3H": "\u00b3H", "3He(trit)": "\u00b3He(trit)", "SF6": "SF\u2086",
          "14C": "\u00b9\u2074C"}
SUP = {"0": "\u2070", "1": "\u00b9", "2": "\u00b2", "3": "\u00b3",
       "4": "\u2074", "5": "\u2075", "6": "\u2076", "7": "\u2077",
       "8": "\u2078", "9": "\u2079", "-": "\u207b"}


def sig3(x):
    """Three significant figures, with a power-of-ten form below 1e-3 --
    matching the precision pass applied to the other SI tables."""
    if not np.isfinite(x):
        return "\u2014"
    if x != 0 and abs(x) < 1e-3:
        e = int(np.floor(np.log10(abs(x))))
        m = x / 10 ** e
        ms = ("%g" % round(m, 2))
        return "%s \u00d7 10%s" % (ms, "".join(SUP[c] for c in str(e)))
    s = "%.3g" % x
    return s


def compute():
    lo = pd.read_csv(LOTO)
    cen = lo.censored_true.fillna(True).astype(bool)
    defined = lo.loc[~cen, "information_ratio_true"].dropna()
    sub = lo.loc[~cen & (lo.information_ratio_true < 1.0)].dropna(
        subset=["information_ratio_true"]).copy()
    sub = sub.sort_values("information_ratio_true")
    return sub, len(defined)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--write-docx", action="store_true",
                    help="also rewrite SI Table S10 and its caption")
    A = ap.parse_args()

    sub, n_def = compute()
    n14 = int((sub.omitted_tracer == "14C").sum())
    print("defined ratios          %d" % n_def)
    print("below unity             %d" % len(sub))
    print("minimum                 %.4f" % sub.information_ratio_true.min())
    print("from omitting 14C       %d" % n14)
    print("by tracer               %s"
          % sub.omitted_tracer.value_counts().to_dict())

    out = pd.DataFrame({
        "SampleID": sub.SampleID,
        "Aq_class": sub.Aq_class,
        "omitted_tracer": sub.omitted_tracer,
        "w_full": sub.w_full_true,
        "w_reduced": sub.w_reduced_true,
        "information_ratio": sub.information_ratio_true,
        "delta_tau1_rel": sub.delta_tau1_rel,
    })
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    out.to_csv(OUT, index=False)
    print("\nwrote %s (%d rows)" % (OUT, len(out)))

    if not A.write_docx:
        print("\n(rerun with --write-docx to update SI Table S10)")
        return

    import copy
    import docx
    d = docx.Document(SI)
    tbl = next((t for t in d.tables
                if "Tracer withheld" in [c.text.strip()
                                         for c in t.rows[0].cells]), None)
    if tbl is None:
        sys.exit("ABORT sub-unity table not found")

    def setc(cell, txt):
        p = cell.paragraphs[0]
        if not p.runs:
            p.add_run("")
        p.runs[0].text = txt
        for r in p.runs[1:]:
            r.text = ""

    proto = tbl.rows[1]._tr
    while len(tbl.rows) - 1 < len(sub):
        proto.addnext(copy.deepcopy(proto))
    while len(tbl.rows) - 1 > len(sub):
        last = tbl.rows[-1]._tr
        last.getparent().remove(last)

    for i, r in enumerate(sub.itertuples(), start=1):
        vals = [r.SampleID, r.Aq_class, PRETTY.get(r.omitted_tracer,
                                                   r.omitted_tracer),
                sig3(r.w_full_true), sig3(r.w_reduced_true),
                sig3(r.information_ratio_true), sig3(r.delta_tau1_rel)]
        for j, v in enumerate(vals):
            setc(tbl.rows[i].cells[j], v)
        print("  row %2d  %s" % (i, " | ".join(vals)))

    old = ("Of the 59 finite ratios, 13 fall below unity, with a minimum of "
           "0.816, because")
    new = ("Of the %d finite ratios, %d fall below unity, with a minimum of "
           "%.3f, because" % (n_def, len(sub),
                              sub.information_ratio_true.min()))
    old2 = "Ten of the thirteen involve \u00b9\u2074C"
    new2 = "Eight of the twelve involve \u00b9\u2074C"
    hits = 0
    for p in d.paragraphs:
        for a, b in ((old, new), (old2, new2)):
            if a in p.text:
                for rr in p.runs:
                    if a in rr.text:
                        rr.text = rr.text.replace(a, b)
                        hits += 1
                        break
    if hits != 2:
        sys.exit("ABORT caption edits matched %d of 2" % hits)
    print("  ok  caption: 13/0.816/ten -> %d/%.3f/eight"
          % (len(sub), sub.information_ratio_true.min()))
    d.save(SI)
    print("saved SI")


if __name__ == "__main__":
    main()
